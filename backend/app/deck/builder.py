"""Build a .pptx deck that presents a generated dashboard.

Driven by the same dashboard spec the dashboard itself renders from, so the
deck always describes the dashboard that actually exists.

Charts are native PowerPoint chart objects rather than screenshots: no headless
browser is needed, and the team can edit them. Their series are **scaffolding**
with an explicit label saying so — the live figures come from the running
dashboard's own API, which this process cannot reach, and inventing numbers and
presenting them as measurements would be worse than leaving the shape to fill in.
"""

from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches, Pt

# 16:9, the shape every projector expects.
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(0x14, 0x16, 0x1C)
MUTED = RGBColor(0x4A, 0x4F, 0x5C)

# Accents match kit/theme.css, so the deck looks like the dashboard it presents.
ACCENTS: dict[str, RGBColor] = {
    "indigo": RGBColor(0x43, 0x38, 0xCA),
    "teal": RGBColor(0x0F, 0x76, 0x6E),
    "amber": RGBColor(0xB4, 0x53, 0x09),
    "rose": RGBColor(0xBE, 0x12, 0x3C),
    "violet": RGBColor(0x6D, 0x28, 0xD9),
    "graphite": RGBColor(0x33, 0x41, 0x55),
}

# The validated categorical palette, in the kit's fixed order.
SERIES_RGB = [
    RGBColor(0x2A, 0x78, 0xD6),
    RGBColor(0xEB, 0x68, 0x34),
    RGBColor(0x1B, 0xAF, 0x7A),
    RGBColor(0xED, 0xA1, 0x00),
    RGBColor(0xE8, 0x7B, 0xA4),
    RGBColor(0x00, 0x83, 0x00),
    RGBColor(0x4A, 0x3A, 0xA7),
    RGBColor(0xE3, 0x49, 0x48),
]

# Module-level so it is not a function call in a default argument (ruff B008).
BODY_TOP = Inches(1.6)

SCAFFOLD_NOTE = "Chart shape only — replace the series with live figures from the dashboard."


def _text_box(slide: Any, left: Emu, top: Emu, width: Emu, height: Emu) -> Any:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def _style(run: Any, size: int, color: RGBColor, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Segoe UI"


def _blank(presentation: Presentation) -> Any:
    """Layout 6 is the blank one in the default template."""
    return presentation.slides.add_slide(presentation.slide_layouts[6])


def _heading(slide: Any, text: str, accent: RGBColor) -> None:
    frame = _text_box(slide, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    run = frame.paragraphs[0].add_run()
    run.text = text
    _style(run, 30, accent, bold=True)


def _bullets(slide: Any, items: list[str], top: Emu = BODY_TOP) -> None:
    frame = _text_box(slide, Inches(0.9), top, Inches(11.5), Inches(5.2))
    for index, item in enumerate(items or ["(not specified)"]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = f"•  {item}"
        _style(run, 16, INK)
        paragraph.space_after = Pt(10)


def _as_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(v) for v in value if str(v).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def _title_slide(presentation: Presentation, spec: dict[str, Any], accent: RGBColor) -> None:
    slide = _blank(presentation)
    frame = _text_box(slide, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6))
    run = frame.paragraphs[0].add_run()
    run.text = spec.get("title") or "Dashboard"
    _style(run, 48, accent, bold=True)

    subtitle = spec.get("subtitle")
    if subtitle:
        paragraph = frame.add_paragraph()
        run = paragraph.add_run()
        run.text = str(subtitle)
        _style(run, 20, MUTED)

    footer = _text_box(slide, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.5))
    run = footer.paragraphs[0].add_run()
    run.text = f"Layout: {spec.get('layout', 'kpi-overview')}"
    _style(run, 12, MUTED)


def _chart_slide(presentation: Presentation, widget: dict[str, Any], accent: RGBColor) -> None:
    """A native, editable chart matching the widget's shape."""
    slide = _blank(presentation)
    _heading(slide, str(widget.get("title") or "Chart"), accent)

    series = [s for s in widget.get("series", []) if isinstance(s, dict)]
    is_bar = widget.get("kind") == "bar"

    data = CategoryChartData()
    data.categories = (
        ["North", "South", "East", "West"] if is_bar else ["Wk 1", "Wk 2", "Wk 3", "Wk 4", "Wk 5"]
    )
    shapes = [
        (42, 38, 51, 46, 58),
        (28, 33, 30, 39, 36),
        (19, 22, 26, 24, 31),
        (12, 15, 13, 18, 17),
        (9, 11, 10, 14, 13),
        (7, 8, 9, 10, 11),
        (5, 6, 7, 7, 8),
        (3, 4, 5, 5, 6),
    ]
    for index, entry in enumerate(series[: len(shapes)] or [{"label": "Value"}]):
        values = shapes[index][: len(data.categories)]
        data.add_series(str(entry.get("label") or f"Series {index + 1}"), values)

    chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if is_bar else XL_CHART_TYPE.LINE_MARKERS
    graphic = slide.shapes.add_chart(
        chart_type, Inches(0.9), Inches(1.6), Inches(11.5), Inches(4.5), data
    ).chart

    chart_series = list(graphic.series)

    # Identity is never colour alone: two or more series always get a legend.
    if len(chart_series) >= 2:
        graphic.has_legend = True
        graphic.legend.position = XL_LEGEND_POSITION.BOTTOM
        graphic.legend.include_in_layout = False
    else:
        graphic.has_legend = False

    for index, plot_series in enumerate(chart_series):
        colour = SERIES_RGB[min(index, len(SERIES_RGB) - 1)]
        if is_bar:
            plot_series.format.fill.solid()
            plot_series.format.fill.fore_color.rgb = colour
        else:
            plot_series.format.line.color.rgb = colour
            plot_series.format.line.width = Pt(2.25)

    note = _text_box(slide, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.5))
    run = note.paragraphs[0].add_run()
    run.text = f"{SCAFFOLD_NOTE}  Source: {widget.get('endpoint', '')}"
    _style(run, 11, MUTED)


def _widget_label(widget: dict[str, Any]) -> str:
    name = widget.get("title") or widget.get("label") or "(untitled)"
    return f"{name}  —  {widget.get('kind', '?')}, from {widget.get('endpoint', '')}"


def build_deck(spec: dict[str, Any], requirements: dict[str, Any]) -> bytes:
    """Return a .pptx presenting the dashboard described by ``spec``."""
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H

    accent = ACCENTS.get(str(spec.get("theme") or "indigo"), ACCENTS["indigo"])

    _title_slide(presentation, spec, accent)

    problem = requirements.get("problem")
    slide = _blank(presentation)
    _heading(slide, "The problem", accent)
    _bullets(slide, _as_list(problem) or ["(not specified)"])

    slide = _blank(presentation)
    _heading(slide, "Who it is for", accent)
    _bullets(slide, _as_list(requirements.get("targetUsers")))

    slide = _blank(presentation)
    _heading(slide, "How it is used", accent)
    _bullets(slide, _as_list(requirements.get("mainWorkflow")))

    widgets = [w for w in spec.get("widgets", []) if isinstance(w, dict)]

    slide = _blank(presentation)
    _heading(slide, "What the dashboard shows", accent)
    _bullets(slide, [_widget_label(w) for w in widgets])

    for widget in widgets:
        if widget.get("kind") in ("line", "bar"):
            _chart_slide(presentation, widget, accent)

    slide = _blank(presentation)
    _heading(slide, "What success looks like", accent)
    _bullets(slide, _as_list(requirements.get("successCriteria")))

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
